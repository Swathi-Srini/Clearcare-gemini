from pathlib import Path
import json
import math
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'CLEARCARE_AMD_ACT_III_SUBMISSION'
VIDEO = OUT / 'videos'
WORK = ROOT / 'CLEARCARE_DEMO_AUTOMATION' / 'presentation_work'
OUT.mkdir(exist_ok=True)
VIDEO.mkdir(exist_ok=True)
WORK.mkdir(exist_ok=True)

NAVY = '#09293C'
TEAL = '#2E7C72'
MINT = '#DDF2EB'
MINT2 = '#B7E5D9'
INK = '#213E45'
MUTED = '#667D7C'
PAPER = '#F4F8F6'
WHITE = '#FFFFFF'
AMBER = '#B87527'
PALE_AMBER = '#FFF1D9'
BLUE = '#386784'
PALE_BLUE = '#EAF2F7'
CORAL = '#D96B5F'

def rgb(hex_color):
    h = hex_color.lstrip('#')
    return RGBColor(*(int(h[i:i+2], 16) for i in (0, 2, 4)))

def font_path(bold=False):
    choices = [
        Path('C:/Windows/Fonts/Manrope-Bold.ttf' if bold else 'C:/Windows/Fonts/Manrope-Regular.ttf'),
        Path('C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf'),
    ]
    return next((str(p) for p in choices if p.exists()), None)

def cover_image(path):
    W, H = 1920, 1080
    im = Image.new('RGB', (W, H), NAVY)
    px = im.load()
    c1 = (8, 35, 54); c2 = (17, 82, 92)
    for x in range(W):
        t = x / (W - 1)
        col = tuple(round(c1[k] * (1-t) + c2[k] * t) for k in range(3))
        for y in range(H): px[x, y] = col
    d = ImageDraw.Draw(im, 'RGBA')
    # soft atmospheric rings
    for inset, alpha in [(0, 9), (35, 12), (72, 15)]:
        d.ellipse((1220+inset, 100+inset, 1870-inset, 750-inset), outline=(170, 245, 225, alpha+35), width=2)
    # brand symbol
    d.rounded_rectangle((142, 135, 206, 199), radius=19, fill=(91, 175, 155, 255))
    d.line((174, 150, 174, 184), fill=(255, 255, 255, 255), width=8)
    d.line((157, 167, 191, 167), fill=(255, 255, 255, 255), width=8)
    regular, bold = font_path(False), font_path(True)
    fbrand = ImageFont.truetype(bold, 46) if bold else ImageFont.load_default()
    ftitle = ImageFont.truetype(bold, 83) if bold else ImageFont.load_default()
    fsub = ImageFont.truetype(regular, 35) if regular else ImageFont.load_default()
    fsmall = ImageFont.truetype(bold, 21) if bold else ImageFont.load_default()
    ffoot = ImageFont.truetype(regular, 20) if regular else ImageFont.load_default()
    d.text((230, 135), 'clearcare.', font=fbrand, fill='#F5FBF8')
    d.text((146, 345), 'The next steps,', font=ftitle, fill='#F7FAF8')
    d.text((146, 440), 'made clear.', font=ftitle, fill='#83D9C7')
    d.rounded_rectangle((151, 565, 250, 573), radius=4, fill=CORAL)
    d.text((148, 608), 'Understand discharge instructions.', font=fsub, fill='#D9EAE6')
    d.text((148, 657), 'Feel more ready for the days ahead.', font=fsub, fill='#D9EAE6')
    d.rounded_rectangle((148, 770, 818, 829), radius=20, fill=(255,255,255,22), outline=(179,231,218,90), width=2)
    d.text((174, 788), 'AMD DEVELOPER CLOUD  ·  TOKEN FACTORY', font=fsmall, fill='#BDEDE2')
    d.text((148, 994), 'Synthetic demo  ·  Source-grounded  ·  Not medical advice', font=ffoot, fill='#B5C9C8')

    # Illustrative care plan card on the right
    d.rounded_rectangle((1025, 236, 1768, 947), radius=36, fill=(244,250,247,248), outline=(186,224,215,255), width=2)
    d.text((1074, 282), 'YOUR CARE PLAN', font=fsmall, fill=TEAL)
    d.text((1074, 321), 'Today at home', font=ImageFont.truetype(bold, 40) if bold else fbrand, fill=INK)
    rows = [
        ('✓', 'Take a short walk', 'As you feel comfortable', MINT, TEAL),
        ('!', 'Confirm a dose', 'Ask your care team first', PALE_AMBER, AMBER),
        ('▦', 'Follow-up visit', 'Confirmed · 10:30 AM', PALE_BLUE, BLUE),
    ]
    y = 412
    for icon, title, detail, bg, fg in rows:
        d.rounded_rectangle((1068,y,1725,y+116), radius=19, fill='#FFFFFF', outline='#DDE9E4', width=2)
        d.rounded_rectangle((1090,y+25,1137,y+72), radius=13, fill=bg)
        if icon == '✓':
            d.line((1103,y+51,1113,y+61,1128,y+40), fill=fg, width=4)
        elif icon == '▦':
            d.rounded_rectangle((1102,y+36,1127,y+61), radius=3, outline=fg, width=3)
            d.line((1103,y+43,1126,y+43), fill=fg, width=2)
            d.line((1109,y+44,1109,y+60), fill=fg, width=2)
            d.line((1118,y+44,1118,y+60), fill=fg, width=2)
        else:
            d.text((1108,y+30), icon, font=fsmall, fill=fg)
        d.text((1160,y+25), title, font=ImageFont.truetype(bold, 27) if bold else fbrand, fill=INK)
        d.text((1160,y+65), detail, font=ImageFont.truetype(regular, 21) if regular else ffoot, fill=MUTED)
        y += 137
    # proof/evidence ribbon
    d.rounded_rectangle((1068, 826, 1725, 922), radius=18, fill='#EAF5F1')
    d.rectangle((1068,826,1075,922), fill=TEAL)
    d.text((1093, 844), 'ORIGINAL WORDS STAY IN VIEW', font=fsmall, fill=TEAL)
    d.text((1093, 879), 'Understand first. Review when unclear.', font=ImageFont.truetype(regular, 19) if regular else ffoot, fill=MUTED)
    im.save(path, optimize=True)


SLIDES = [
    {
        'title': 'Recovery deserves understanding.',
        'eyebrow': 'CLEARCARE  ·  HEALTH & WELLBEING',
        'body': 'A hospital discharge can leave patients and family caregivers with a lot to remember. ClearCare turns the written next steps into a plan that is easier to follow.',
        'narration': 'After a hospital stay, patients and family caregivers may be sent home with pages of instructions. ClearCare helps make those next steps easier to understand and organize, while keeping the care team’s original words in view.',
        'visual': 'home', 'duration': 12
    },
    {
        'title': 'Start with the document.',
        'eyebrow': '01  ·  READ',
        'body': 'Select a text-based discharge PDF. The demo reads it in the browser and keeps the selected file in the current session.',
        'narration': 'The workflow starts with a discharge PDF. ClearCare reads text-based documents in the browser. The demonstration uses synthetic patient details, and uploaded files are not saved to an account.',
        'visual': 'document', 'duration': 11
    },
    {
        'title': 'Keep each instruction grounded.',
        'eyebrow': '02  ·  SHOW THE SOURCE',
        'body': 'Medication directions and care steps stay paired with the exact words found in the document, so families can compare the summary with its source.',
        'narration': 'Each extracted instruction stays connected to its source wording. Families can see the original line alongside the organized plan, instead of having to trust a summary on its own.',
        'visual': 'evidence', 'duration': 11
    },
    {
        'title': 'Explain simply, when you ask.',
        'eyebrow': '03  ·  AMD TOKEN FACTORY',
        'body': 'A person chooses Explain simply on one instruction. ClearCare sends that selected text to the backend AMD service for a short plain-language explanation.',
        'narration': 'The core AI action is explicit. Nothing is sent just because the page opened. When a person chooses Explain simply, ClearCare sends only the selected instruction to the backend, which calls AMD Developer Cloud Token Factory using DeepSeek V4 Flash.',
        'visual': 'amd', 'duration': 14
    },
    {
        'title': 'Unclear stays unclear.',
        'eyebrow': '04  ·  SAFETY BY DESIGN',
        'body': 'Missing units, ambiguous doses, conflicting directions, and high-risk medicines are marked for review. ClearCare does not guess or mark them confirmed.',
        'narration': 'When an instruction is ambiguous, missing a unit, conflicting, or involves a high-risk medicine, ClearCare routes it for review. It does not silently choose a direction, confirm an item, or schedule an incomplete appointment.',
        'visual': 'review', 'duration': 13
    },
    {
        'title': 'A plan that fits real life.',
        'eyebrow': '05  ·  CHECKLIST & CALENDAR',
        'body': 'Every task has its own checkbox. Confirmed appointments can be shown on the calendar and exported as an .ics file; items needing review stay out.',
        'narration': 'The care plan uses a straightforward checklist, with separate checkboxes for separate tasks. Confirmed follow-ups can appear on the calendar and be exported as an ICS file. Items that still need review are not eligible for scheduling.',
        'visual': 'calendar', 'duration': 12
    },
    {
        'title': 'Make the plan easier to use.',
        'eyebrow': '06  ·  LANGUAGE & VOICE',
        'body': 'English, Spanish, and Hindi are available in Settings. Today · Hear my day adds a readable transcript, browser voice controls, and document questions.',
        'narration': 'ClearCare includes English, Spanish, and Hindi settings, plus a Today view with a readable transcript and browser-based voice controls. People can type or speak a question and edit the text before sending it.',
        'visual': 'language', 'duration': 12
    },
    {
        'title': 'AI that knows its limits.',
        'eyebrow': '07  ·  HOW AMD FITS',
        'body': 'AMD Token Factory is a hosted API integration. The API key stays on the backend. ClearCare does not use ROCm or run the model on the patient’s device.',
        'narration': 'The AI runs through AMD Developer Cloud Token Factory, not locally through ROCm. The key remains on the backend. ClearCare is a prototype organizer and explainer, not a clinician, and the care team remains the source of medical direction.',
        'visual': 'architecture', 'duration': 13
    },
    {
        'title': 'ClearCare',
        'eyebrow': 'AMD DEVELOPER HACKATHON  ·  ACT III',
        'body': 'Clear next steps for the days after discharge. Synthetic demo · source-grounded · human review where needed.',
        'narration': 'ClearCare makes the next steps after discharge easier to see, understand, and organize. This is a synthetic-data prototype for the AMD Developer Hackathon. It does not diagnose or prescribe, and it never replaces guidance from the care team.',
        'visual': 'close', 'duration': 13
    }
]

def set_shape_fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()

def box(slide, x, y, w, h, color, radius=True, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_fill(shape, color)
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(1)
    else: shape.line.fill.background()
    return shape

def text(slide, value, x, y, w, h, size, color=INK, bold=False, font='Aptos', align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(0); tf.margin_right = Inches(0); tf.margin_top = Inches(0); tf.margin_bottom = Inches(0); tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align; p.space_before = Pt(0); p.space_after = Pt(0)
    r = p.add_run(); r.text = value; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
    return shape

def small_label(slide, value, x, y, w=5.5, color=TEAL): text(slide, value.upper(), x, y, w, .24, 9, color, True)

def icon_card(slide, x, y, w, h, title, detail, icon, state='normal'):
    bg = PALE_AMBER if state == 'review' else WHITE
    border = '#E8D5B2' if state == 'review' else '#DDE9E4'
    box(slide,x,y,w,h,bg,True,border)
    box(slide,x+.18,y+.19,.43,.43, PALE_AMBER if state=='review' else MINT, True)
    text(slide,icon,x+.18,y+.18,.43,.43,8 if len(icon)>1 else 13,AMBER if state=='review' else TEAL,True,align=PP_ALIGN.CENTER)
    text(slide,title,x+.75,y+.15,w-.9,.3,13,INK,True)
    text(slide,detail,x+.75,y+.5,w-.9,.35,9,MUTED)

def visual(slide, kind):
    x,y,w,h = 8.35,1.66,4.2,4.9
    box(slide,x,y,w,h,WHITE,True,'#DDE9E4')
    if kind == 'home':
        small_label(slide,'A calmer path home',x+.28,y+.28,3)
        text(slide,'Your next steps',x+.28,y+.64,3.5,.46,23,INK,True)
        icon_card(slide,x+.25,y+1.35,3.7,.94,'Care plan','A clear place to start','OK')
        icon_card(slide,x+.25,y+2.43,3.7,.94,'Needs review','Questions stay visible','!','review')
        icon_card(slide,x+.25,y+3.51,3.7,.94,'Calendar','Confirmed dates in view','DATE')
    elif kind == 'document':
        small_label(slide,'Synthetic discharge summary',x+.3,y+.3,3.4)
        box(slide,x+.3,y+.75,3.6,3.72,'#FCFDFC',True,'#E4ECE9')
        text(slide,'DISCHARGE INSTRUCTIONS',x+.56,y+1.02,3.1,.28,9,TEAL,True)
        text(slide,'Medication',x+.56,y+1.56,3,.27,12,INK,True)
        text(slide,'Acetaminophen 500 mg every 6 hours as needed for pain.',x+.56,y+1.92,3,.7,11,MUTED)
        box(slide,x+.56,y+2.83,2.9,.025,'#DDE9E4',False)
        text(slide,'Follow-up',x+.56,y+3.02,3,.27,12,INK,True)
        text(slide,'Primary care · Oct 20 · 10:30 AM',x+.56,y+3.39,3,.55,10,MUTED)
        box(slide,x+.57,y+4.0,1.6,.24,MINT,True)
        text(slide,'READ IN BROWSER',x+.67,y+4.01,1.4,.21,7,TEAL,True)
    elif kind == 'evidence':
        small_label(slide,'Medication instruction',x+.3,y+.3,3)
        text(slide,'Acetaminophen · 500 mg',x+.3,y+.82,3.6,.38,17,INK,True)
        text(slide,'Every 6 hours as needed for pain',x+.3,y+1.25,3.55,.4,11,MUTED)
        box(slide,x+.3,y+2.0,3.6,1.72,'#F2F7F5',True)
        box(slide,x+.3,y+2.0,.05,1.72,TEAL,False)
        small_label(slide,'Original words',x+.55,y+2.22,3)
        text(slide,'“Take acetaminophen 500 mg every 6 hours as needed for pain.”',x+.55,y+2.62,3.05,.8,12,INK)
        box(slide,x+.3,y+4.0,2.1,.53,TEAL,True)
        text(slide,'AI  ·  Explain simply',x+.43,y+4.08,1.85,.32,9,WHITE,True,align=PP_ALIGN.CENTER)
    elif kind == 'amd':
        small_label(slide,'One click · one selected instruction',x+.3,y+.3,3.7)
        box(slide,x+.35,y+1.1,1.3,1.06,MINT,True)
        text(slide,'Care text',x+.45,y+1.38,1.1,.3,12,TEAL,True,align=PP_ALIGN.CENTER)
        box(slide,x+2.55,y+1.1,1.3,1.06,PALE_BLUE,True)
        text(slide,'AMD API',x+2.65,y+1.38,1.1,.3,12,BLUE,True,align=PP_ALIGN.CENTER)
        text(slide,'→',x+1.82,y+1.42,.5,.3,24,TEAL,True,align=PP_ALIGN.CENTER)
        text(slide,'Selected text only',x+.35,y+2.5,1.3,.5,9,MUTED,align=PP_ALIGN.CENTER)
        text(slide,'DeepSeek-V4-Flash',x+2.45,y+2.5,1.5,.5,9,MUTED,align=PP_ALIGN.CENTER)
        box(slide,x+.55,y+3.35,3.1,1.08,'#F7FAF9',True,'#DDE9E4')
        text(slide,'Plain words · source stays visible',x+.76,y+3.65,2.68,.42,11,INK,True,align=PP_ALIGN.CENTER)
        text(slide,'Backend key · no request on page load',x+.48,y+4.55,3.25,.26,8,TEAL,True,align=PP_ALIGN.CENTER)
    elif kind == 'review':
        small_label(slide,'Safety review',x+.3,y+.3,3)
        text(slide,'Needs a second look',x+.3,y+.68,3.5,.45,18,INK,True)
        icon_card(slide,x+.25,y+1.38,3.7,.94,'Dose unit missing','Ask the care team to confirm','!','review')
        icon_card(slide,x+.25,y+2.46,3.7,.94,'Directions conflict','Both source lines stay visible','!','review')
        icon_card(slide,x+.25,y+3.54,3.7,.94,'High-risk medicine','Never auto-confirm or schedule','!','review')
    elif kind == 'calendar':
        small_label(slide,'A simple checklist',x+.3,y+.3,3)
        icon_card(slide,x+.25,y+.85,3.7,.8,'Buy medication: A','A separate reminder','1')
        icon_card(slide,x+.25,y+1.78,3.7,.8,'Buy medication: B','Its own checkbox','2')
        box(slide,x+.25,y+2.8,3.7,1.52,'#F4F8F6',True,'#DDE9E4')
        text(slide,'OCTOBER 2026',x+.48,y+3.0,3.2,.27,9,TEAL,True)
        for i,n in enumerate(['M','T','W','T','F','S','S']): text(slide,n,x+.46+i*.45,y+3.38,.28,.22,8,MUTED,True,align=PP_ALIGN.CENTER)
        for i,n in enumerate(range(19,26)):
            box(slide,x+.43+i*.45,y+3.7,.33,.34, MINT if n==20 else WHITE,True,'#E4ECE9')
            text(slide,str(n),x+.43+i*.45,y+3.73,.33,.25,8,TEAL if n==20 else INK,n==20,align=PP_ALIGN.CENTER)
        text(slide,'Only complete, confirmed visits export',x+.35,y+4.55,3.5,.26,8,MUTED,align=PP_ALIGN.CENTER)
    elif kind == 'language':
        small_label(slide,'Built-in care plan languages',x+.3,y+.3,3.5)
        for i,(native,name) in enumerate([('EN','English'),('ES','Español'),('हि','हिन्दी')]):
            yy=y+.92+i*.8
            box(slide,x+.28,yy,3.64,.62,MINT if i==0 else '#F8FAF9',True,'#E0EAE6')
            box(slide,x+.45,yy+.11,.4,.4,TEAL if i==0 else PALE_BLUE,True)
            text(slide,native,x+.45,yy+.14,.4,.27,9,WHITE if i==0 else BLUE,True,align=PP_ALIGN.CENTER)
            text(slide,name,x+1.0,yy+.13,2.5,.3,12,INK,True)
        box(slide,x+.28,y+3.55,3.64,.9,'#F6FAF8',True,'#E0EAE6')
        text(slide,'Read the transcript.\nPlay, pause, or stop.',x+.5,y+3.7,3.1,.55,11,MUTED)
        text(slide,'Voice availability depends on the device.',x+.3,y+4.62,3.6,.22,8,MUTED)
    elif kind == 'architecture':
        small_label(slide,'The request path',x+.3,y+.3,3)
        labels=[('01','Patient chooses','one explanation'),('02','ClearCare backend','holds the secret key'),('03','AMD Token Factory','DeepSeek-V4-Flash'),('04','Verified result','shown with source')]
        for i,(num,title,sub) in enumerate(labels):
            yy=y+.72+i*.91
            box(slide,x+.3,yy,3.6,.72, '#F6FAF8' if i!=2 else PALE_BLUE,True,'#E0EAE6')
            text(slide,num,x+.48,yy+.17,.4,.26,9,TEAL if i!=2 else BLUE,True)
            text(slide,title,x+1.02,yy+.09,2.55,.27,11,INK,True)
            text(slide,sub,x+1.02,yy+.39,2.55,.22,8,MUTED)
        text(slide,'Not ROCm · no local model runtime',x+.4,y+4.52,3.5,.28,9,BLUE,True,align=PP_ALIGN.CENTER)
    elif kind == 'close':
        small_label(slide,'ClearCare',x+.3,y+.35,3)
        text(slide,'One step at a time.',x+.3,y+1.08,3.5,.9,26,INK,True)
        text(slide,'Your care team remains your guide.',x+.3,y+2.05,3.5,.55,14,TEAL,True)
        box(slide,x+.3,y+3.18,3.6,1.05,MINT,True)
        text(slide,'SYNTHETIC DEMO',x+.55,y+3.52,3.1,.31,11,TEAL,True,align=PP_ALIGN.CENTER)

def build_deck(cover):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]
    for idx,item in enumerate(SLIDES):
        s=prs.slides.add_slide(blank)
        bg=s.background.fill; bg.solid(); bg.fore_color.rgb=rgb(PAPER if idx not in (0,8) else NAVY)
        if idx==0:
            s.shapes.add_picture(str(cover),0,0,width=prs.slide_width,height=prs.slide_height)
        elif idx==8:
            s.shapes.add_picture(str(cover),0,0,width=prs.slide_width,height=prs.slide_height)
            box(s,.55,6.9,6.3,.36,NAVY,True)
            text(s,'CLEARCARE · SYNTHETIC PROTOTYPE',.73,6.96,5.9,.21,8,'#B5C9C8',True)
        else:
            box(s,.55,.47,.11,.36,TEAL,False)
            small_label(s,item['eyebrow'],.83,.5,7)
            text(s,item['title'],.83,.94,7.05,1.05,27,INK,True)
            text(s,item['body'],.83,2.18,6.75,1.65,14,MUTED,False,valign=MSO_ANCHOR.TOP)
            visual(s,item['visual'])
            box(s,.83,5.93,6.85,.01,'#DDE9E4',False)
            text(s,'ClearCare  ·  Synthetic demo  ·  Not medical advice',.83,6.13,6.8,.28,8,MUTED)
            text(s,f'{idx:02d} / {len(SLIDES)-1:02d}',11.85,6.92,.9,.23,8,MUTED,True,align=PP_ALIGN.RIGHT)
        # Speaker notes are stored as narration source even if a machine voice is unavailable.
        try: s.notes_slide.notes_text_frame.text = item['narration']
        except Exception: pass
    path=OUT/'ClearCare_Hackathon_Presentation.pptx'; prs.save(path)
    (WORK/'narration.json').write_text(json.dumps([{'id':f'slide-{i+1:02d}','text':s['narration'],'duration':s['duration']} for i,s in enumerate(SLIDES)],ensure_ascii=False,indent=2),encoding='utf-8')
    return path

def main():
    cover=OUT/'cover.png'
    cover_image(cover)
    deck=build_deck(cover)
    print(json.dumps({'cover':str(cover),'coverSize':Image.open(cover).size,'presentation':str(deck),'slides':len(SLIDES),'videoFolder':str(VIDEO),'narrationSource':str(WORK/'narration.json')},indent=2))

if __name__=='__main__': main()
